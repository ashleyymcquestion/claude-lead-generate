"""PI doctor lead workbook + CSV. Contact fields hold ONLY what a cited search result showed; blank = not retrieved."""
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

V = 'https://www.dr-virella.com/Personal-Injury-Spine-Surgeon-Expert-SoCal.html'
NY = 'https://nyspine.com/personal-injury'
AS = 'https://asappaindocs.com/about-advanced-spine-and-pain'
GA = 'https://markets.financialcontent.com/stocks/article/pressadvantage-2025-3-20-georgia-spine-and-orthopaedics-launches-car-accident-recovery-program-with-top-car-accident-doctor-team'
TX = 'https://versustexas.com/blog/treatment-after-an-accident/'
# (establishment, doctors, specialty, state, city, phone, email, address, pi_evidence, fit, source, verified, notes)
R = [
 ('Accident Clinic AZ LLC','not retrieved','Accident injury clinic','AZ','Phoenix (3 locations)','','','','Directory text: lien-based care for MVA victims; 3 locations; plans multi-state expansion','Strong: growing, lien-based','https://www.doccafe.com/company/137544/accident-clinic-az-llc','partial','Phone/address not in search results'),
 ('Advanced Spine and Pain (ASAP)','Abram Burgher, Todd Turley, Jarrett Leathem; also Roy O\'Neil DO, Mark Kallgren MD','Pain / ortho / neuro / TBI / chiro / PT','AZ','Phoenix HQ; ~15 clinics + surgery centers','480-573-0130 (main); 520-653-0146 (Tucson)','','','Describes a dedicated PI team working with law firms from onboarding through settlement','Strong: multi-specialty PI','https://www.asappaindocs.com/node/436','partial','Doctor names/phones from search snippets of asappaindocs.com; confirm'),
 ('New York Spine Institute','Surnames only: Demoura, Macagno, Desai, Roberts, Post','Spine / ortho / neurosurgery','NY','Westbury (main), Bronx, Brooklyn, Elmhurst, Manhattan, Long Island','1-888-444-6974 (NYSI)','','','States it treats PI and workers\' comp on a lien basis',"Strong: surgical lien volume",NY,'partial','Attorney liaison per page: Kourtney Castelli, ext 1031 (staff, not a physician). Full doctor names not retrieved'),
 ('Dr. Anthony Virella, MD, FACS','Anthony Virella, MD, FACS','Neurological spine surgery','CA','Los Angeles / Ventura County','805-449-0088','','','States he works with LA PI attorneys on a lien basis','Strong: lien surgical cases',V,'partial','Phone from page title "Consultation & Appointments"'),
 ('Georgia Spine & Orthopaedics','not retrieved','Orthopedic / neuro / pain','GA','Roswell (11650 Alpharetta Hwy #100), Tucker, Cumming','678-730-0948','','11650 Alpharetta Hwy #100, Roswell, GA 30076','Car Accident Recovery Program (press release)','Possible: confirm program is lien-based',GA,'partial','Lien/LOP not stated in sources'),
 ('Atlanta Personal Injury Doctors','not retrieved','Chiro / PT / pain / ortho network','GA','Atlanta metro','','','','Press release: lien-basis treatment for uninsured accident patients','Possible: lien network','https://www.newswire.com/atlanta-personal-injury-doctors/121150','partial','No contact details in results'),
 ('RC Chiropractic & Personal Injury Centers','not retrieved','Chiropractic','GA','Duluth / Buckhead','','','','Accident/PI focus; lien not confirmed','Possible','https://next-www.thumbtack.com/ga/duluth/holistic-medicine/rc-chiropractic-personal-injury-centers/service/479862708305977346','unverified',''),
 ('Spine & Pain Institute of Texas','not retrieved','Pain mgmt / chiropractic','TX','DFW metroplex (many locations)','','','','Search result: dedicated LOP coordination program','Multi-site LOP receivables',TX,'partial',''),
 ('Texas Pain & Injury Chiropractic','not retrieved','Chiropractic','TX','DFW / San Antonio / Austin','','','','Search result: accepts letters of protection','LOP receivables',TX,'partial',''),
 ('Texas Injury and Rehab Solutions','not retrieved','Chiro / ortho spine / pain / imaging','TX','Plano / Waxahachie','','','','Search result: accepts letters of protection','LOP receivables',TX,'partial',''),
 ('Accident Centers of Texas','not retrieved','Chiropractic / rehab','TX','Dallas / Fort Worth / Austin (claimed)','','','','Named in one search summary as an LOP network; a follow-up search could NOT confirm it exists as a multi-site chain','Unconfirmed: verify existence first',TX,'unverified','Downgraded from earlier list'),
 ('Accident & Rehab Center','1 chiropractor (name not retrieved)','Chiropractic','TX','Dallas (3456 Webb Chapel Ext, 75220)','214-902-8868','','3456 Webb Chapel Ext, Dallas, TX 75220','Accident-focused chiropractic; LOP/lien not confirmed','Possible','https://practicefinder.newpatientsinc.com/chiro/accident-rehab-center-dallas-tx/','unverified','Phone from directory listing'),
 ('The Injury Docs','not retrieved','Chiropractic + MD','FL','Orlando / Central FL','','','','Auto accident / PI network; PIP accepted; LOP not confirmed','Possible','https://www.zoominfo.com/c/the-injury-docs/370459468','unverified',''),
 ('The Chiropractic Clinics of South Florida','not retrieved','Chiropractic','FL','Miami-Dade / Broward / West Palm Beach','','','','Accident injury clinic; LOP not confirmed','Possible multi-site','https://dorisaveslives.org/about/safety-partners/our-safety-partners/the-chiropractic-clinics-of-south-florida.html','unverified',''),
 ('Elite SpineCare','not retrieved','Chiropractic','FL','Kissimmee','','','','Auto accident focus; PIP and attorney paperwork','Small; confirm LOP','https://lantern.llc/b/elite-spinecare-kissimmee','unverified',''),
 ('Saadat Spine','not retrieved','Orthopedic spine','CA','LA / Orange County','','','','Surfaced for PI spine care; lien terms not confirmed','Possible','https://www.wboc.com/online_features/press_releases/leading-spine-surgeon-in-los-angeles-saadat-spine-responds-to-rising-demand-for-advanced-orthopedic/article_42a39a94-a5e6-50ac-aafc-d6afd63ec472.html','unverified',''),
 ('Dr. Hormoz Zahiri','Hormoz Zahiri (USC clinical professor)','Orthopedic spine','CA','Los Angeles','','','','Listed in a PI / workers\' comp physician network (Becker\'s)','Possible','https://www.beckersspine.com/?p=12262','unverified',''),
 ('The Neck and Back Clinics','not retrieved','Chiropractic / rehab','NV','Las Vegas','','','','Described as covering personal injury and auto accident cases','Possible','https://www.cbinsights.com/company/the-neck-and-back-clinics','unverified','New this round'),
 ('Nevada Orthopedic & Spine Center','not retrieved (fellowship-trained spine surgeons)','Spine surgery','NV','Las Vegas / Henderson','702-258-3773','','','No PI/lien evidence in results; included as a spine-surgery prospect only','Weak: qualify PI/lien volume','https://nevadaorthopedic.com/contents/additional-services','unverified','New this round'),
]
H = ['Establishment','Doctors (as found)','Specialty','State','City / area','Phone','Email','Address','PI / lien evidence','Funding fit (inferred)','Source URL','Verified','Notes']
with open('06_pi_doctor_leads.csv','w',newline='') as f:
    w = csv.writer(f); w.writerow(H); w.writerows(R)

wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'PI Doctor Leads'
ws.append(['#'] + H + ['Email verified?','DNC scrubbed?','Date called','Spoke with','Outcome','Next step'])
for i, r in enumerate(R, 1): ws.append([i] + list(r) + ['']*6)
for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F3864'); c.alignment = Alignment(wrap_text=True, vertical='center')
for i, wd in enumerate([4,30,34,24,6,28,26,26,30,40,26,40,10,38,12,11,12,16,22,22], 1): ws.column_dimensions[get_column_letter(i)].width = wd
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
    v = str(row[12].value)
    row[12].fill = PatternFill('solid', fgColor='FFEB9C' if v == 'partial' else 'F8CBAD')
ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions
n = wb.create_sheet('Read me')
for line in ['PI doctor leads (USA), generated 2026-10-01 from web search results only.',
 'Blank Phone/Email/Doctors cells mean NOT RETRIEVED, not "none". Practice sites could not be opened (network proxy block); no emails were found in any result.',
 'Phones are business main lines from search snippets/directories. Verify on each practice site, and scrub against DNC lists before calling.',
 'No source shows a practice is seeking funding; "Funding fit" is inferred from lien/LOP treatment. Qualify on the first call.',
 'Doctor names are only those a result printed. Several are surnames only.']: n.append([line])
n.column_dimensions['A'].width = 140
wb.save('PFD_PI_Doctor_Leads.xlsx'); print(len(R), 'rows;', sum(1 for r in R if r[5]), 'with phone;', sum(1 for r in R if r[6]), 'with email')
