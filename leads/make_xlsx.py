"""Detailed workbook matching PFD_Call_Sheet.pdf (same firms, same order)."""
import csv, glob, re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

EXCLUDE = ('ocreia', 'real estate investors association', 'windfall advisors')
rows = []
for r in csv.DictReader(open('firms_with_phone.csv')):
    if any(x in r['Firm'].lower() for x in EXCLUDE): continue
    rows.append(dict(firm=r['Firm'], cat=r['Category'], city=r['City'].split('(')[0].strip(), state='CA', region='Orange County, CA',
        website=r['Website'], phone=r['Main phone (verify before calling)'], contact=r['Contact name'], title=r['Contact title'],
        email='', conf=r['Contact confidence'], src=r['Contact source URL'], notes=r['Cautions']))
for f in sorted(glob.glob('national_*.csv')):
    region = f[9:-4].replace('_', ' ').title()
    for r in csv.DictReader(open(f)):
        ph = (r.get('main_phone') or '').strip()
        if not ph or not re.search(r'\d{3}\D*\d{4}', ph): continue
        rows.append(dict(firm=r['firm'], cat=r['category'], city=r['city'], state=r['state'], region=region, website=r.get('website', ''),
            phone=ph, contact=r.get('contact_name', ''), title=r.get('contact_title', ''), email=r.get('business_email', ''),
            conf=r.get('confidence', ''), src=r.get('source_url', ''), notes=r.get('notes', '')))
seen, uniq = set(), []
for r in rows:
    k = re.sub(r'\D', '', r['phone'])[-10:]
    if k in seen: continue
    seen.add(k); uniq.append(r)
uniq.sort(key=lambda r: (r['state'] != 'CA', r['state'], r['firm'].lower()))
for r in uniq: r['state'] = r['state'].strip()

wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Call Sheet Firms'
hdr = ['#', 'Firm', 'Type', 'City', 'State', 'Region', 'Main phone', 'Contact name', 'Contact title', 'Business email', 'Website',
       'Confidence', 'Source URL', 'Notes', 'Phone verified on firm site?', 'DNC scrubbed?', 'Date called', 'Spoke with', 'Outcome', 'Next step', 'Follow-up date']
ws.append(hdr)
for i, r in enumerate(uniq, 1):
    ws.append([i, r['firm'], r['cat'], r['city'], r['state'], r['region'], r['phone'], r['contact'], r['title'], r['email'], r['website'],
               r['conf'], r['src'], r['notes'], '', '', '', '', '', '', ''])
for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F3864'); c.alignment = Alignment(wrap_text=True, vertical='center')
for i, w in enumerate([5, 30, 22, 16, 7, 18, 16, 22, 20, 24, 28, 11, 34, 40, 14, 11, 12, 16, 22, 22, 13], 1): ws.column_dimensions[get_column_letter(i)].width = w
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions
fills = {'high': 'C6EFCE', 'medium': 'FFEB9C', 'low': 'F8CBAD'}
for row in ws.iter_rows(min_row=2):
    v = str(row[11].value).lower()
    for k, col in fills.items():
        if v.startswith(k): row[11].fill = PatternFill('solid', fgColor=col)

s = wb.create_sheet('Summary')
s.append(['Breakdown', 'Firms'])
from collections import Counter
s.append(['TOTAL', len(uniq)])
s.append([]); s.append(['By region', 'Firms'])
for k, v in Counter(r['region'] for r in uniq).most_common(): s.append([k, v])
s.append([]); s.append(['By state', 'Firms'])
for k, v in sorted(Counter(r['state'] for r in uniq).items()): s.append([k, v])
s.append([]); s.append(['By confidence', 'Firms'])
for k, v in Counter((r['conf'] or 'unrated').lower() for r in uniq).most_common(): s.append([k, v])
s.append([]); s.append(['Notes'])
for n in ['Business main lines only; numbers come from search listings and are unverified.',
          'Scrub every number against National, California and internal DNC lists before calling.',
          'Use attorney-approved scripts; never promise or guarantee returns.']: s.append([n])
for c in s[1]: c.font = Font(bold=True)
s.column_dimensions['A'].width = 40
wb.save('PFD_Call_Sheet_Details.xlsx'); print(len(uniq), 'firms')
