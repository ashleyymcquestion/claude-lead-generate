"""Workbook of lien-doctor directories/networks (lead SOURCES) + practices named in searches. Blank = not retrieved."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

D = [  # name, type, coverage, website, phone, what it offers, how to use for leads, source, verified
 ('Doctors on Liens, Inc.','Lien-doctor network / directory','CA, NV, AZ','https://doctorsonliens.com','','Since 1993; lists med/legal doctors by specialty (orthopedist, urgent care, etc.) who treat PI and workers\' comp on lien; printed map (Spring 2026)','Browse specialty pages and printed map for provider names; contact the company about listing/partnering','https://doctorsonliens.com/specialists/Orthopedist','partial'),
 ('Power Liens','Lien-doctor directory','CA (claims other US areas)','https://powerliens.com','','Describes itself as the largest directory of PI doctors on liens in CA; attorney-facing; has Northern CA blog','Browse provider listings by specialty/city; practices listed here already take liens','https://powerliens.com/attorneys','partial'),
 ('California Lien Providers (CLP)','Lien-provider directory','CA','https://www.lienproviders.com','(800) 466-6824','Directory of specialists willing to treat on lien; staff with 25+ years in PI treatment facilitation','Call/browse for specialist lists; ask about provider onboarding','https://www.lienproviders.com/','partial'),
 ('Injury Institute','Medical network on liens','CA','https://injuryinstitute.com','','Claims to be the largest lien medical network in CA: physicians, surgeons, chiropractors','Network members are lien-treating practices; approach as potential funding partner','https://injuryinstitute.com/','partial'),
 ('SoCal Injury Liens','Lien doctor network','CA, AZ, NV','https://www.socalinjuryliens.com','','Lien doctor network for PI attorneys','Same as above','https://www.socalinjuryliens.com/','partial'),
 ('Liens Studios','Marketing/tech platform for PI providers','US (multiple states)','https://www.liensstudios.com','','Publishes guides on PI medical provider networks and credentialing; serves medical providers','Likely has provider clients; confirm what it offers','https://www.liensstudios.com/practice-areas/medical-providers','unverified'),
 ('Empower MH trusted provider network','Provider network (chiropractic listing page)','unknown','https://www.empowermh.co','','Has a "trusted provider network" chiropractic page','Confirm geography and whether lien-based','https://www.empowermh.co/trusted-provider-network/providers/chiropractic','unverified'),
 ('Becker\'s ASC: online physicians directory for attorneys','Trade article on a physician directory','AZ, CA, FL, IL (per search)','https://www.beckersasc.com/?p=21178','','Article about an online physicians directory for attorneys adding pain management; names not retrieved','Read article to identify the directory operator','https://www.beckersasc.com/?p=21178','unverified'),
 ('Ventura County Trial Lawyers Association (VCTLA)','Trial-lawyer association','CA (Ventura)','https://www.vcba.org','','Surfaced in search; no medical-provider directory found in results','Ask about sponsor/vendor lists','https://www.vcba.org/wp-content/uploads/2026/04/042826-VCTLA.pdf','unverified'),
 ('Consumer Attorneys of California','Trial-lawyer association','CA','','','Named in results; no provider directory found','Ask about sponsor/vendor lists','','unverified'),
 ('Illinois Chiropractic Society','State chiropractic association','IL','https://ilchiro.org','','Has health care lien resources for members','Request member directory; filter for PI','https://ilchiro.org/health-care-liens/','unverified'),
 ('California Chiropractic Association','State chiropractic association','CA','','','Provides lien forms for PI recovery','Request member directory; filter for PI','','unverified'),
]
P = [ # practices surfaced by these searches
 ('Arrowhead Clinic Chiropractor','Chiropractic','GA','Brunswick','','','Press release: expanded partnership network with PI attorneys for no-cost accident treatment','https://www.barchart.com/story/news/449774/arrowhead-clinic-chiropractor-brunswick-announces-expanded-partnership-network-with-personal-injury-attorneys-for-no-cost-accident-treatment','partial'),
]
wb = openpyxl.Workbook(); a = wb.active; a.title = 'Lien directories & networks'
ha = ['#','Name','Type','Coverage','Website','Phone','What it offers','How to use for leads','Source URL','Verified']
a.append(ha)
for i, r in enumerate(D, 1): a.append([i] + list(r))
b = wb.create_sheet('Practices named')
hb = ['#','Practice','Specialty','State','City','Phone','Email','PI / lien evidence','Source URL','Verified']
b.append(hb)
for i, r in enumerate(P, 1): b.append([i] + list(r))
for ws, widths in ((a,[4,30,28,18,30,16,50,44,40,10]),(b,[4,30,16,7,14,14,22,44,40,10])):
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F3864'); c.alignment = Alignment(wrap_text=True, vertical='center')
    for i, wd in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = wd
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
        row[-1].fill = PatternFill('solid', fgColor='FFEB9C' if row[-1].value == 'partial' else 'F8CBAD')
    ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions
n = wb.create_sheet('Read me')
for t in ['Generated 2026-10-01 from web search results only. The directory sites (doctorsonliens.com, powerliens.com, lienproviders.com) are blocked in this environment, so the individual doctors they list could NOT be copied.',
          'Sheet 1 lists the directories/networks themselves as lead SOURCES. Blank cells = not retrieved.',
          'Next step: open each directory in a browser; export or copy its listings, or contact the operator. Respect each site\'s terms of use (many prohibit scraping) and DNC/TCPA rules.',
          'No state trial-lawyer or chiropractic association directory of lien-treating providers was found; those entries are only places to request member lists.']: n.append([t])
n.column_dimensions['A'].width = 150
wb.save('PFD_Lien_Directories.xlsx'); print(len(D), len(P))
